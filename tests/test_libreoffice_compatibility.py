from __future__ import annotations

import os
import shutil
import subprocess
from pathlib import Path

import pytest
from docx import Document

from docxrender import DocxRenderer


def _write_template(path: Path) -> None:
    document = Document()
    document.add_paragraph("{{ report_title }}")
    document.add_paragraph("{{ body_anchor }}")
    document.save(str(path))


@pytest.mark.skipif(
    os.environ.get("DOCXRENDER_TEST_LIBREOFFICE") != "1",
    reason="set DOCXRENDER_TEST_LIBREOFFICE=1 to run the real conversion probe",
)
def test_markdown_table_converts_with_real_libreoffice(tmp_path: Path) -> None:
    executable = shutil.which("libreoffice")
    if executable is None:
        pytest.fail("DOCXRENDER_TEST_LIBREOFFICE=1 but libreoffice is unavailable")

    template = tmp_path / "template.docx"
    report = tmp_path / "report.docx"
    output_directory = tmp_path / "pdf"
    profile_directory = tmp_path / "profile"
    output_directory.mkdir()
    _write_template(template)

    (
        DocxRenderer()
        .with_template(
            file_template=template,
            context={"report_title": "LibreOffice table-grid compatibility"},
        )
        .write_docx(
            file_out_docx=report,
            markdown_body=(
                "| A | B | C |\n"
                "| --- | --- | --- |\n"
                "| one | two | three |\n"
                "| four | five | six |\n"
            ),
            dir_base=tmp_path,
        )
    )

    completed = subprocess.run(
        [
            executable,
            "--headless",
            f"-env:UserInstallation={profile_directory.resolve().as_uri()}",
            "--convert-to",
            "pdf",
            "--outdir",
            str(output_directory),
            str(report),
        ],
        check=False,
        capture_output=True,
        text=True,
        timeout=30,
    )
    output_pdf = output_directory / "report.pdf"
    assert completed.returncode == 0, completed.stdout + completed.stderr
    assert output_pdf.is_file(), completed.stdout + completed.stderr
    assert output_pdf.read_bytes().startswith(b"%PDF-")
